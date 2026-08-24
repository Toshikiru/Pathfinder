"""One-off script: fills docs/Project_Proposal.docx from the instructor's
blank template (../Project Proposal.docx) with Pathfinder's Checkpoint 1
proposal content, preserving the template's own formatting/styles."""

import docx
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Project Proposal.docx"
OUT = ROOT / "docs" / "Project_Proposal.docx"


def set_para_text(p, text: str):
    """Replace a paragraph's visible text, keeping the first run's formatting."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_para_after_label(p, addition: str):
    """Keep the first (bold label) run, replace/clear the rest, append addition."""
    if len(p.runs) >= 2:
        p.runs[1].text = " " + addition
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(" " + addition)


def main():
    d = docx.Document(str(SRC))

    # Instructor line (template has a <w:br/> before this run - the .text
    # setter clears child elements, so re-insert the line break after).
    instructor_run = d.paragraphs[1].runs[-1]
    instructor_run.text = "Instructor: Jessie A. Melendres"
    br = instructor_run._r.makeelement(qn("w:br"), {})
    t = instructor_run._r.find(qn("w:t"))
    instructor_run._r.insert(list(instructor_run._r).index(t), br)

    # Chosen theme checklist (paras 5-11)
    set_para_text(d.paragraphs[5], "☑ Document/Knowledge-Base Q&A Assistant  (SELECTED)")
    set_para_text(d.paragraphs[6], "☐ Source Code Analyzer/Reviewer")
    set_para_text(d.paragraphs[7], "☐ Customer Support / FAQ Bot")
    set_para_text(d.paragraphs[8], "☐ Academic Research Assistant")
    set_para_text(d.paragraphs[9], "☐ Local Business / Community Assistant")
    set_para_text(d.paragraphs[10], "☐ Personal/Domain Notes Assistant")
    set_para_text(d.paragraphs[11], "☐ Others (please specify): N/A")

    d.paragraphs[13].runs[-1].text = (
        "Repository Link: https://github.com/<your-team>/pathfinder-rag-advisor "
        "(create before submission; all checkpoint commits live here per Section 7 of the Guide)"
    )

    # 2. Problem Statement
    set_para_text(d.paragraphs[15], (
        "At Talibon Polytechnic College (TPC), BSIS students plan their course loads and "
        "check their degree progress manually: cross-referencing a printed curriculum "
        "checklist, remembering which subjects are prerequisites for which, and asking the "
        "Program Head or Registrar's Office in person for anything not obvious from the "
        "checklist (unit caps under probation, GWA requirements for Latin honors, whether a "
        "grade of INC blocks enrollment in a follow-on course, etc.). This does not scale: "
        "the Program Head and Registrar's Office serve every student in the program with "
        "limited office hours, so students commonly enroll in courses whose prerequisites "
        "they have not completed, miscount their remaining units toward graduation, or simply "
        "do not find out about a policy (e.g. the probation unit cap) until it already affects "
        "them. Existing non-AI solutions - a static PDF/printed checklist and word-of-mouth "
        "from advisers - cannot answer a specific, personalized question (\"Can I take IS301 "
        "next term if I haven't passed IS205 yet?\") and cannot be queried outside office hours. "
        "Pathfinder addresses this by grounding a conversational assistant in TPC's actual "
        "BSIS curriculum structure and academic policies, so students get an accurate, "
        "cited answer to course-planning and degree-progress questions on demand."
    ))

    # 3. Target Users
    set_para_text(d.paragraphs[17], (
        "Primary users: current BSIS students at TPC, especially those who are enrolling, "
        "re-enrolling after a failed/dropped course, or nearing graduation and need to verify "
        "remaining requirements. Secondary users: incoming/prospective BSIS students comparing "
        "program requirements before enrolling, and the Program Head/Registrar's Office staff, "
        "who can point students to the assistant for routine prerequisite and policy questions "
        "instead of answering them individually. For primary users, the tool replaces manually "
        "cross-referencing a paper checklist and waiting for office hours with an instant, "
        "always-available, source-cited answer; for staff, it reduces repetitive front-desk "
        "questions so office hours can focus on cases that genuinely need a human (petitions, "
        "waivers, appeals)."
    ))

    # 4. Dataset Source(s)
    set_para_text(d.paragraphs[19], (
        "The knowledge base combines two kinds of sources, documented in full in "
        "docs/Data_Sources.md: (1) real content scraped from TPC's public website "
        "(tpc.edu.ph) - program descriptions, program leadership, admission requirements, "
        "enrollment process, and student services info; and (2) a simulated BSIS curriculum "
        "checklist and academic policy set, since TPC's actual registrar prospectus is not "
        "published online. The simulated documents are modeled on CHED Memorandum Order "
        "No. 25, s. 2015 (the real national policy standard for BSIS program design in the "
        "Philippines) and are clearly labeled [SIMULATED DATA] in-file and in the README."
    ))
    set_para_after_label(d.paragraphs[21], "Plain text (.txt), scraped from HTML and generated from structured Python data.")
    set_para_after_label(d.paragraphs[22], "tpc.edu.ph (public website, real) + a CHED CMO No. 25 s.2015 - based simulated curriculum/policy set (clearly labeled).")
    set_para_after_label(d.paragraphs[23], "30 documents total: 16 real scraped pages + 9 curriculum-checklist documents + 3 course-description catalogs + 2 policy documents.")

    # 5. Success Criteria
    set_para_text(d.paragraphs[25], (
        "By the end of the semester, Pathfinder is successful if it meets all of the following:"
    ))
    set_para_after_label(d.paragraphs[27], "the system correctly cites the source document for at least 90% of course-planning/policy queries in the Checkpoint 3 demo set.")
    set_para_after_label(d.paragraphs[28], "retrieval latency (query embedding + similarity search) stays under 2 seconds on the full corpus.")
    set_para_after_label(d.paragraphs[29], "the assistant correctly answers at least 5 realistic multi-part advising questions (e.g. \"Can I take IS301 if I haven't passed IS205?\") without hallucinating a course, code, or policy that isn't in the knowledge base.")

    # Insert 2 more criteria after para 29 to reach 5 total by cloning its XML.
    import copy
    new_p1 = copy.deepcopy(d.paragraphs[29]._p)
    new_p2 = copy.deepcopy(d.paragraphs[29]._p)
    d.paragraphs[29]._p.addnext(new_p2)
    d.paragraphs[29]._p.addnext(new_p1)

    # Re-fetch paragraph objects now that new ones are inserted (indices shift by +2)
    p30 = d.paragraphs[30]
    p31 = d.paragraphs[31]
    set_para_text(p30, "Refusal Correctness: the assistant correctly declines or flags at least 1 out-of-scope/unanswerable query per demo session instead of guessing.")
    set_para_text(p31, "Coverage: all four BSIS year levels (1-4) and both simulated policy documents are each retrieved correctly at least once across the demo query set.")

    table = d.tables[0]
    table.cell(1, 0).text = "[Team Member 1 - Your Name]"
    table.cell(2, 0).text = "[Team Member 2]"
    table.cell(3, 0).text = "[Team Member 3]"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(OUT))
    print(f"Saved filled proposal to {OUT}")


if __name__ == "__main__":
    main()
